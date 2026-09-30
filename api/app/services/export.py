"""Dataset export (CSV / JSON / XLSX) and run-to-run diffs."""

import csv
import io
import json
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..models import Provenance, Record
from ..pipeline.merge import entity_key

TRUST_COLUMNS = ["trust", "confidence", "sources"]


def _rows(intent: dict, records: list[Record], prov: list[Provenance], sources: dict[int, str]):
    fields = [f["name"] for f in intent.get("fields", [])]
    urls = defaultdict(list)
    for p in prov:
        u = sources.get(p.source_id)
        if u and u not in urls[p.record_id]:
            urls[p.record_id].append(u)
    rows = []
    for r in records:
        row = {f: r.data.get(f) for f in fields}
        row |= {"trust": r.status, "confidence": round(r.confidence, 2), "sources": " | ".join(urls[r.id])}
        rows.append(row)
    return fields, rows


def build(fmt: str, intent: dict, records: list[Record], prov: list[Provenance], sources: dict[int, str]) -> bytes:
    fields, rows = _rows(intent, records, prov, sources)
    header = fields + TRUST_COLUMNS

    if fmt == "json":
        cells = defaultdict(lambda: defaultdict(list))
        for p in prov:
            cells[p.record_id][p.field].append({"value": p.value, "quote": p.quote, "url": sources.get(p.source_id)})
        out = [{**row, "evidence": cells[r.id]} for row, r in zip(rows, records)]
        return json.dumps({"entity": intent.get("entity"), "rows": out}, ensure_ascii=False, indent=2).encode()

    if fmt == "csv":
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=header)
        w.writeheader()
        w.writerows(rows)
        return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel reads UTF-8

    wb = Workbook()
    ws = wb.active
    ws.title = "Dataset"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="5B5BD6")
    fills = {"verified": "E3F5EC", "conflict": "FBF0E1"}
    for row in rows:
        ws.append([row.get(h) for h in header])
        color = fills.get(row["trust"])
        if color:
            ws.cell(ws.max_row, header.index("trust") + 1).fill = PatternFill("solid", fgColor=color)
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r.get(h) or "")) for r in rows[:200]])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def diff(id_field: str, before: list[Record], after: list[Record]) -> dict:
    """+new / -removed / ~changed between two runs, matched by entity name."""
    old = {entity_key(r.data.get(id_field, "")): r for r in before}
    new = {entity_key(r.data.get(id_field, "")): r for r in after}
    added = [new[k].data.get(id_field) for k in new if k not in old]
    removed = [old[k].data.get(id_field) for k in old if k not in new]
    changed = []
    for k in new.keys() & old.keys():
        fields = sorted(f for f in set(new[k].data) | set(old[k].data)
                        if str(new[k].data.get(f)) != str(old[k].data.get(f)))
        if fields:
            changed.append({"name": new[k].data.get(id_field), "record_id": new[k].id, "fields": fields})
    return {"added": added, "removed": removed, "changed": changed, "unchanged": len(new) - len(added) - len(changed)}
